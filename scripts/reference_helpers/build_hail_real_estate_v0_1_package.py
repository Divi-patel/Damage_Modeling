"""Deterministic builder for the noncanonical hail × real estate model-v0.1 / docs-r1 package.

One pathway (hail_impact) × the three Hazard reference profiles r1 × two priced failure units: the roof cover
(Haag hail-size thresholds × Hazus roof cost shares) and a Tier-4 walls-and-windows screening unit that Hazard must
acknowledge (decision log D-HRE-01 to D-HRE-16). Exposed equipment and contents are withheld, not zero.

Reads the governed CSVs in docs/cells/hail_real_estate/proposed/ (the ordinates come from
derive_hail_real_estate_curves.py; the NRI check from hail_real_estate_nri_check.py) and writes:
    hail_real_estate__model_v0_1__docs_r1__curve_artifact.json   (bundle v3, embedded capability)
    hail_real_estate__model_v0_1__docs_r1__capability.json       (identical standalone capability)
    known_answer_tests_hail_real_estate__model_v0_1__docs_r1.json
    damage_curve_records_hail_real_estate__model_v0_1__docs_r1.xlsx

Known-answer expectations are computed here from the ordinates CSV with a separate, minimal interpolation, not by
calling the evaluator. Requires openpyxl for the workbook.

Usage (from the repository root):
    python scripts/reference_helpers/build_hail_real_estate_v0_1_package.py
"""
from __future__ import annotations

import csv
import datetime
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CELL_ID = "hail_real_estate"
CELL = f"docs/cells/{CELL_ID}/proposed"
TAG = f"{CELL_ID}__model_v0_1__docs_r1"
PATHWAY = "hail_impact"
AXIS = "mesh_diameter_mm"
ROOF, WALLS = "HRE_ROOF_COVER", "HRE_WALLS_AND_WINDOWS"
EQUIP, CONTENTS = "HRE_EXPOSED_EQUIPMENT", "HRE_BUILDING_CONTENTS"
ACK = "screening_assumption_set_id"
ACK_ID = "HRE_T4_WALLS_WINDOWS_SCREENING_2026_10_18"
LO, HI = 0.0, 203.2
MESH75_FLOOR = 38.98  # MESH75 of the 25.4 mm raw-MESH event threshold (Hazard)
EXPECTED_RECORDS = 6
BUILDINGS = {
    "RE_RES_SFD_1S_SLAB": {"profile": "re_house_1s_wood_slab", "roof": "asphalt shingles (laminated), 10 years",
                           "flags": ["ROOF_ASPHALT_SHINGLE_LAMINATED_ASSUMED_NOT_IMPACT_RATED"]},
    "RE_IND_WAREHOUSE": {"profile": "re_warehouse_1s_steel_slab", "roof": "standing-seam steel, 10 years",
                         "flags": ["METAL_ROOF_FUNCTIONAL_DAMAGE_ONLY_COSMETIC_DENTING_EXCLUDED",
                                   "WAREHOUSE_ROOF_COST_SHARE_JUDGMENT"]},
    "RE_COM_OFFICE_3S_SLAB": {"profile": "re_office_3s_slab", "roof": "TPO membrane on insulated metal deck, 10 years",
                              "flags": ["TPO_ATTACHMENT_AND_COVER_BOARD_NOT_STATED"]},
}
UNIT_FLAGS = {
    ROOF: ["ROOF_COVER_HAIL_THRESHOLD_EVIDENCE", "WHOLE_ROOF_COVER_REPLACED_WHEN_DAMAGED",
           "ROOF_COVER_SHARE_FROM_HAZUS_COST_RATIOS_OR_JUDGMENT"],
    WALLS: ["TIER4_SCREENING_WALLS_AND_WINDOWS", "NO_PUBLIC_WALL_OR_WINDOW_HAIL_THRESHOLD",
            "HAIL_FACING_SHARE_JUDGMENT"],
}
ALWAYS = ["NONCANONICAL_PROPOSAL", "SCREENING_MATERIAL_THRESHOLD_CURVES", "COMPONENT_FAILURE_UNITS_OWN_VALUE_DENOMINATOR",
          "AXIS_IS_HAZARD_MESH75_READ_AS_HAILSTONE_DIAMETER", "THRESHOLD_BAND_READ_AS_90_PERCENT_RANGE",
          "COSMETIC_DENTING_EXCLUDED", "CONSEQUENTIAL_INTERIOR_LEAKAGE_NOT_PRICED",
          "CUMULATIVE_SUB_SEVERE_HAIL_NOT_MODELLED", "REFERENCE_BUILDING_ASSUMPTIONS_NOT_OBSERVED",
          "NOT_FIELD_OR_CLAIMS_CALIBRATED", "PLATFORM_CODES_PROPOSED_NOT_ISSUED"]
WITHHELD = {
    EQUIP: ["NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT", "FUNCTIONAL_EQUIPMENT_HAIL_EVIDENCE_NOT_FOUND"],
    CONTENTS: ["NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT", "CONTENTS_DEFERRED_SEPARATE_FAILURE_UNIT"],
}


def rows(name: str) -> list[dict[str, str]]:
    with open(ROOT / CELL / name, newline="") as fh:
        return list(csv.DictReader(fh))


def ordinates() -> dict[tuple[str, str], list[list[float]]]:
    pts: dict[tuple[str, str], list[list[float]]] = {}
    for r in rows(f"CURVE_ORDINATES_{TAG}.csv"):
        pts.setdefault((r["reference_building_id"], r["failure_unit_id"]), []).append(
            [float(r["mesh_diameter_mm"]), float(r["damage_ratio"])])
    return {k: sorted(v) for k, v in pts.items()}


def curve_ids() -> dict[tuple[str, str], str]:
    return {(r["reference_building_id"], r["failure_unit_id"]): r["curve_id"] for r in rows(f"CURVE_ORDINATES_{TAG}.csv")}


def parts() -> dict[tuple[str, str], list[dict[str, str]]]:
    out: dict[tuple[str, str], list[dict[str, str]]] = {}
    for r in rows(f"COMPONENT_FRAGILITY_TABLE_{TAG}.csv"):
        out.setdefault((r["reference_building_id"], r["failure_unit_id"]), []).append(r)
    return out


def unit_share(building: str, unit: str) -> float:
    return round(sum(float(p["part_value_fraction_of_structure_rcn"]) for p in parts()[(building, unit)]), 6)


def interp(points: list[list[float]], x: float) -> float:
    i = max(j for j, p in enumerate(points) if p[0] <= x)
    if points[i][0] == x:
        return points[i][1]
    (x0, y0), (x1, y1) = points[i], points[i + 1]
    return y0 + (y1 - y0) * (x - x0) / (x1 - x0)


def curve_records(pts) -> list[dict]:
    ids, prt = curve_ids(), parts()
    out = []
    for unit in (ROOF, WALLS):
        for building, spec in BUILDINGS.items():
            p = pts[(building, unit)]
            ps = prt[(building, unit)]
            match = {"reference_building_id": building}
            if unit == WALLS:
                match[ACK] = ACK_ID
            refs = []
            for part in ps:
                refs += [f"{s}#{part['part_id']}" for s in part["source_ids"].split(";")]
            out.append({
                "curve_id": ids[(building, unit)], "pathway_id": PATHWAY, "failure_unit_id": unit,
                "curve_form": "piecewise_linear", "x_axis": AXIS,
                "y_axis": "failure_unit_damage_ratio", "parameters": {"points": p},
                "valid_range": [p[0][0], p[-1][0]],
                "interpolation_policy": "linear_between_source_knots",
                "extrapolation_policy": ("knots every 1 mm from 0 to 203 mm plus 203.2 mm (tabulated from the lognormal "
                                         "part-replacement expression); above 203.2 mm withholds; negative rejected; no clamp"),
                "selector_match": match,
                "source_parameter_refs": list(dict.fromkeys(refs)),
                "metadata_flags": spec["flags"] + UNIT_FLAGS[unit],
            })
    return out


def capability() -> dict:
    return {
        "schema_version": "capability_declaration.v3", "cell_id": CELL_ID,
        "canonical_runtime_artifact": False,
        "pathway_capabilities": [{
            "pathway_id": PATHWAY,
            "failure_unit_scalar_dr": "conditional",
            "scenario_loss_given_value_basis": "withheld",
            "curve_intrinsic_spread": "not_carried",
            "populated_emit_modes": ["scalar_mean"],
            "conditions": [
                "reference_building_id must be one of the three Hazard reference profiles r1 buildings",
                "mesh_diameter_mm is Hazard's MESH75 hail diameter (mm) as delivered, read as the largest hailstone at the building",
                "records are valid from 0 to 203.2 mm; above 203.2 withholds; a negative diameter is rejected",
                "each failure unit's damage ratio is a fraction of that unit's own replacement value; the unit's share "
                "of structure replacement cost new is in the value crosswalk and failure_units[].share_of_structure_rcn",
                f"the walls-and-windows unit is Tier-4 screening and answers only with {ACK} = {ACK_ID}",
                "exposed equipment and contents are withheld, never zero",
                "scenario dollars stay withheld until the consumer qualifies the building value",
                "owner acceptance of the scope and selections is required before any canonical use",
            ],
            "limitation_flags": ALWAYS + ["VALUE_SHARES_PARTLY_JUDGMENT", "MULTI_PERIL_SAME_BUILDING_COMBINATION_CONSUMER_OWNED"],
            "withheld_failure_units": [{"failure_unit_id": u, "reason_codes": c} for u, c in WITHHELD.items()],
        }],
        "consumer_annual_metrics": {
            "computation_owner": "downstream_consumer",
            "status_before_promotion": "withheld_noncanonical_proposal",
            "status_after_promotion": "withheld",
            "prerequisites": ["owner acceptance of scope and curve selections",
                              "consumer-qualified replacement cost new for each reference building",
                              "Hazard review of the MESH75 floor and point hail frequency (NRI check)",
                              "exact model/docs/schema/SHA pin"],
            "limitation_flags": ["ANNUAL_AND_TAIL_METRICS_CONSUMER_OWNED", "NRI_CHECK_ABOVE_NRI_EXPECTED_UNDERREPORTING"],
        },
        "cap_binding": {
            "policy": "consumer_enforced_fail_closed", "enforcement_owner": "downstream_consumer",
            "checks_required": ["exact model/docs/schema/SHA pin",
                                "record routing by failure_unit_id and reference_building_id",
                                "unit DR multiplied by the same unit's share of structure value, never by the whole value",
                                "walls-and-windows unit used only with the screening acknowledgement",
                                "all perils on one building in one event together at most 100% of structure value"],
            "action_if_fail": "withhold the affected DR or downstream metric; do not substitute another record or clamp",
        },
        "promotion_gate": {"status": "blocked", "required_before_canonical_use": [
            "owner accepts D-HRE-01 to D-HRE-16",
            "consumer value qualification of each reference budget (shared with the other real-estate cells)",
            "reviewer accepts the MESH75-as-hailstone axis reading and the Tier-4 walls-and-windows unit",
            "independent review of the package"]},
    }


def axis() -> dict:
    return {"id": "HAIL_DIAMETER_MESH_EQUIV", "preferred_input_field": AXIS, "unit": "mm", "valid_range": [LO, HI],
            "source_units_allowed": ["mm"], "permitted_proxy_fields": [],
            "bridge": ("none in Damage; Hazard delivers MESH75 = 15.096 x ((raw MESH / 2.54)^2)^0.206 per native pixel, "
                       "capped at 203.2 mm, and zero damage below the 25.4 mm raw-MESH threshold; the smallest nonzero "
                       f"input is therefore {MESH75_FLOOR} mm"),
            "category_guardrail": "NWS severe-hail labels and report counts are not damage inputs; a cell maximum or averaged MESH is not a point hail size.",
            "extrapolation_policy": "diameter < 0 rejects; 0-203.2 mm straight lines between 1-mm knots; > 203.2 withholds"}


def failure_units() -> list[dict]:
    share = {b: {u: unit_share(b, u) for u in (ROOF, WALLS)} for b in BUILDINGS}
    return [
        {"id": ROOF, "subsystem": "BUILDING", "component": "ROOF_COVER", "treatment": "primary_nonzero",
         "y_axis": "failure_unit_damage_ratio",
         "denominator": "replacement cost of the same building's roof cover (house shingles; warehouse steel panels; office TPO membrane and insulation)",
         "share_of_structure_rcn": {b: share[b][ROOF] for b in BUILDINGS},
         "platform_code_status": "proposed_awaiting_platform_codes",
         "source_atom_note": "Material hail thresholds (Haag) for the whole roof cover; the cover is replaced whole when damaged."},
        {"id": WALLS, "subsystem": "BUILDING", "component": "WALLS_AND_WINDOWS", "treatment": "secondary_conditional",
         "y_axis": "failure_unit_damage_ratio",
         "denominator": "replacement cost of the same building's hail-vulnerable wall cladding and windows (house vinyl siding and windows; warehouse and office windows)",
         "share_of_structure_rcn": {b: share[b][WALLS] for b in BUILDINGS},
         "platform_code_status": "proposed_awaiting_platform_codes",
         "source_atom_note": "Tier-4 judgment sizes, hail-facing shares and unit costs; requires the screening acknowledgement."},
        {"id": EQUIP, "subsystem": "BUILDING", "component": "EXPOSED_EQUIPMENT", "treatment": "withheld",
         "y_axis": "failure_unit_damage_ratio", "denominator": "replacement cost of the outdoor AC condenser / rooftop units (not allocated)",
         "withheld_reason_codes": WITHHELD[EQUIP], "platform_code_status": "proposed_awaiting_platform_codes"},
        {"id": CONTENTS, "subsystem": "BUILDING", "component": "CONTENTS", "treatment": "withheld",
         "y_axis": "failure_unit_damage_ratio", "denominator": "contents value (deferred)",
         "withheld_reason_codes": WITHHELD[CONTENTS], "platform_code_status": "proposed_awaiting_platform_codes"},
    ]


def pathway_block(pts) -> dict:
    return {
        "pathway_id": PATHWAY,
        "hazard_scope": {
            "included": ["hailstone impact on the roof cover (puncture, fracture, rupture, open seams)",
                         "hailstone impact on vinyl siding and window glazing (Tier-4 screening unit)"],
            "excluded": ["cosmetic denting of metal roofs, metal wall panels, overhead doors and gutters",
                         "condenser and rooftop-unit fin denting (exposed equipment withheld)",
                         "rain entry and interior damage after a roof breach; roof deck punctures by giant hail",
                         "wind from the same storm (convective_wind_real_estate)", "flooding",
                         "contents, loss of use, business interruption, land, site works"],
            "deferred_pathways": ["hail_accumulation_or_meltwater_flood"],
            "compound_event_rule": "Hail and wind from one storm on one building together at most 100% of structure value (consumer).",
        },
        "hazard_axis": axis(),
        "reference_profile_binding": {b: s["profile"] for b, s in BUILDINGS.items()},
        "record_routing_fields": ["failure_unit_id", "reference_building_id"],
        "selector_logic": [
            {"field": "reference_building_id", "required": True, "default": None, "allowed": list(BUILDINGS),
             "effect": "Fixed identity of the reference building; picks its roof cover and wall/window parts."},
            {"field": ACK, "required": False, "required_when": f"failure_unit_id = {WALLS}", "default": None,
             "allowed": [ACK_ID],
             "effect": "Acknowledges the Tier-4 walls-and-windows screening assumptions; not needed for the roof cover."},
        ],
        "conditioner_logic": [],
        "exposure_contract": {
            "required_subject_grain": "one reference building at one native radar pixel (Hazard area-weights pixels for larger footprints)",
            "value_rule": "Unit loss = unit DR x unit share of structure replacement cost new x the consumer-qualified budget.",
            "point_hail_rule": "Hazard's per-pixel MESH75 is read as the largest hailstone at the building; a cell maximum is not.",
        },
        "failure_unit_coverage": [
            {"failure_unit_id": ROOF, "status": "conditional_numeric_screening_curve", "reason_codes": []},
            {"failure_unit_id": WALLS, "status": "conditional_synthetic_t4_curve", "reason_codes": []},
            {"failure_unit_id": EQUIP, "status": "withheld_not_zero", "reason_codes": WITHHELD[EQUIP]},
            {"failure_unit_id": CONTENTS, "status": "withheld_not_zero", "reason_codes": WITHHELD[CONTENTS]},
        ],
        "curve_records": curve_records(pts),
    }


def artifact(pts) -> dict:
    return {
        "schema_version": "damage_curve_record_bundle.v3", "schema_status": "proposed_draft",
        "cell_id": CELL_ID, "damage_code_id": "HAIL_REAL_ESTATE_ROOF_THRESHOLD_SCREENING_V0_1",
        "semantic_damage_model_version": "model v0.1", "documentation_revision": "docs r1",
        "lifecycle_state": "scaffold", "promotion_status": "proposed", "review_status": "not_yet_reviewed",
        "model_grade": "screening_material_threshold_component_curves",
        "package_release": "unreleased", "package_baseline": "library v2.5", "package_inclusion_status": "not_included",
        "canonical_runtime_artifact": False,
        "source_dossier": f"{CELL}/{CELL_ID}_curve_derivation_dossier__model_v0_1__docs_r1.md",
        "source_workbook": f"{CELL}/damage_curve_records_{TAG}.xlsx",
        "known_answer_tests": f"{CELL}/known_answer_tests_{TAG}.json",
        "source_register": f"{CELL}/SOURCE_REGISTER_{TAG}.csv",
        "claim_parameter_register": f"{CELL}/CLAIM_PARAMETER_REGISTER_{TAG}.csv",
        "value_crosswalk": f"{CELL}/VALUE_CROSSWALK_{TAG}.csv",
        "failure_units": failure_units(),
        "pathways": [pathway_block(pts)],
        "value_linkage": {
            "curve_denominator": "replacement cost of the same failure unit",
            "unit_share_basis": "share of structure replacement cost new of the reference building (failure_units[].share_of_structure_rcn)",
            "building_level_formula": "building DR = sum over units of unit DR x unit share",
            "scenario_loss_formula": "loss = unit DR x unit share x structure_replacement_cost_new_usd",
            "scenario_loss_status": "withheld_consumer_value_qualification",
            "rejected_denominators": ["whole structure value as a component denominator", "depreciated value",
                                      "market value", "assessed value"],
            "implicit_default_profile": None,
        },
        "parameter_tier_table": rows(f"PARAMETER_TIER_TABLE_{TAG}.csv"),
        "evaluation_contract": {
            "supported_pathway_ids": [PATHWAY], "supported_failure_unit_ids": [ROOF, WALLS],
            "required_selector_fields": ["reference_building_id"],
            "conditionally_required_selector_fields": {ACK: f"failure_unit_id = {WALLS}"},
            "required_conditioner_fields": [],
            "accepted_axis_payloads": [{"pathway_id": PATHWAY, "mode": "hail_diameter_mm", "required_fields": [AXIS]}],
            "metadata_flags_always": ALWAYS,
            "out_of_range_behavior": "above 203.2 mm: withhold; negative: error",
            "reference_evaluator": "scripts/reference_helpers/hail_real_estate_curve_eval.py",
        },
        "emit_contract": {
            "schema_version": "damage_emit.v2",
            "populated_emit_modes_for_this_cell": ["scalar_mean"],
            "supported_output": "conditional scalar DR per failure unit (roof cover; walls and windows)",
            "scenario_loss": "withheld until the consumer qualifies the building value",
            "annual_and_tail_metrics": "withheld; consumer-owned",
            "withheld_units_emit": "explicit null DR plus reason codes; never zero",
        },
        "capability_declaration": capability(),
        "derivation_rationale": {
            "source_spine": "Haag hail-damage threshold sizes by roofing material (AMS 2002 lab and field tables; Professional Roofing 2009 low-slope table; Haag 2010 metal rupture via CLM 2016), with roof-cover cost shares from FEMA Hazus 7.0 subassembly cost ratios (house, office) and judgment (warehouse).",
            "adopted_form": "per part, replacement probability lognormal in hail diameter (onset = 5%, full = 95%); unit DR = value-weighted, hail-facing-share-weighted sum; tabulated every 1 mm",
            "why_not_swiss_radar_curve": "Schmid et al. 2024 is a portfolio average over Swiss tile-roof buildings in a 1 km pixel (0.05% at 60 mm); owner choice D-HRE-01 keeps it as a cross-check.",
            "parameter_tier_mix": "Roof thresholds T3; house and office roof shares T2/T3 (Hazus); warehouse roof share, warehouse full-damage size, band reading and the whole walls-and-windows unit T4.",
            "nri_check": "See dossier §13 and NRI_CHECK; order-of-magnitude only.",
        },
        "legacy_comparison": {"prior_cell_state": "none (new cell)",
                              "changes_from_prior": [],
                              "legacy_arrays_used": "none; the legacy Real Estate_Hail array is rejected (HRE-C018)",
                              "canonical_or_consumer_change": False},
    }


def known_answer_tests(pts) -> dict:
    ids = curve_ids()

    def req(building, unit, d, **extra):
        r = {"pathway_id": PATHWAY, "failure_unit_id": unit, "reference_building_id": building, AXIS: d, **extra}
        if unit == WALLS and ACK not in extra:
            r[ACK] = ACK_ID
        return r

    formula, withheld, errors = [], [], []
    for building in BUILDINGS:
        for unit in (ROOF, WALLS):
            p, cid = pts[(building, unit)], ids[(building, unit)]
            for x in (0, 10, 20, 25, 30, 35, 40, 45, 50, 55, 60, 65, 70, 75, 80, 90, 100, 150, 203):
                formula.append({"test_id": f"{cid}_KNOT_{x}MM", "input": req(building, unit, float(x)),
                                "expected": {"status": "conditional", "curve_id": cid, "failure_unit_damage_ratio": interp(p, float(x))}})
            for x in (25.4, 31.75, MESH75_FLOOR, 44.45, 50.8, 63.5, 76.2, 88.9, 101.6, 203.2):  # 1-4 in; MESH75 floor; cap
                formula.append({"test_id": f"{cid}_AT_{x:g}MM", "input": req(building, unit, x),
                                "expected": {"status": "conditional", "curve_id": cid, "failure_unit_damage_ratio": interp(p, x)}})
            withheld.append({"test_id": f"{cid}_ABOVE_RANGE", "input": req(building, unit, HI + 0.0001),
                             "expected": {"status": "withheld", "reason_code": "ABOVE_SOURCE_VALID_RANGE"}})
        for unit in (EQUIP, CONTENTS):
            withheld.append({"test_id": f"{building}_{unit}_WITHHELD", "input": req(building, unit, 50.8),
                             "expected": {"status": "withheld", "reason_code": "NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT"}})
        formula.append({"test_id": f"{building}_ROOF_ZERO_AT_ZERO", "input": req(building, ROOF, 0.0),
                        "expected": {"status": "conditional", "curve_id": ids[(building, ROOF)], "failure_unit_damage_ratio": 0.0}})
        formula.append({"test_id": f"{building}_ROOF_WITH_ACK_SAME_VALUE", "input": req(building, ROOF, 50.8, **{ACK: ACK_ID}),
                        "expected": {"status": "conditional", "curve_id": ids[(building, ROOF)],
                                     "failure_unit_damage_ratio": interp(pts[(building, ROOF)], 50.8),
                                     "metadata_flags_include": ["ROOF_COVER_HAIL_THRESHOLD_EVIDENCE"]}})
    base = req("RE_RES_SFD_1S_SLAB", ROOF, 50.8)
    drop = lambda k, b=base: {key: v for key, v in b.items() if key != k}  # noqa: E731
    walls = req("RE_RES_SFD_1S_SLAB", WALLS, 50.8)
    for tid, inp, code in [
        ("PATHWAY_REQUIRED", drop("pathway_id"), "PATHWAY_ID_REQUIRED"),
        ("PATHWAY_UNKNOWN", {**base, "pathway_id": "straight_line_convective"}, "PATHWAY_ID_UNKNOWN"),
        ("BUILDING_REQUIRED", drop("reference_building_id"), "REFERENCE_BUILDING_ID_REQUIRED"),
        ("BUILDING_UNKNOWN", {**base, "reference_building_id": "RE_COM_LOWRISE"}, "REFERENCE_BUILDING_ID_UNKNOWN"),
        ("BUILDING_PROFILE_ID_NOT_A_SELECTOR", {**base, "reference_building_id": "re_house_1s_wood_slab"}, "REFERENCE_BUILDING_ID_UNKNOWN"),
        ("AXIS_REQUIRED", drop(AXIS), "AXIS_PAYLOAD_REQUIRED"),
        ("AXIS_INCHES_FIELD_NOT_ACCEPTED", {**drop(AXIS), "hail_diameter_in": 2.0}, "AXIS_PAYLOAD_REQUIRED"),
        ("AXIS_NOT_NUMBER", {**base, AXIS: "50.8"}, "AXIS_VALUE_INVALID"),
        ("AXIS_BOOLEAN", {**base, AXIS: True}, "AXIS_VALUE_INVALID"),
        ("AXIS_NEGATIVE", {**base, AXIS: -1.0}, "AXIS_OUTSIDE_VALID_RANGE"),
        ("UNIT_REQUIRED", drop("failure_unit_id"), "FAILURE_UNIT_ID_REQUIRED"),
        ("UNIT_UNKNOWN", {**base, "failure_unit_id": "HRE_BUILDING_STRUCTURE"}, "FAILURE_UNIT_ID_UNKNOWN"),
        ("WALLS_WITHOUT_ACKNOWLEDGEMENT", drop(ACK, walls), "SCREENING_ASSUMPTION_SET_REQUIRED"),
        ("WALLS_WRONG_ACKNOWLEDGEMENT", {**walls, ACK: "WW_T4_PARTIAL_ELECTRICAL_SCREENING_2026_08_08"}, "SCREENING_ASSUMPTION_SET_UNSUPPORTED"),
    ]:
        errors.append({"test_id": tid, "input": inp, "expected_error_code": code})
    return {
        "schema_version": "known_answer_tests.v3", "cell_id": CELL_ID,
        "semantic_damage_model_version": "model v0.1", "documentation_revision": "docs r1",
        "artifact_schema_version": "damage_curve_record_bundle.v3", "artifact_status": "noncanonical_proposal",
        "absolute_tolerance": 1e-12,
        "equation": {"runtime": "straight lines between the 1-mm knots of the selected record",
                     "axis": "mesh_diameter_mm (Hazard's MESH75 hail diameter, mm)",
                     "ranges": "0 to 203.2 mm evaluates; above 203.2 mm withholds; negative is an error"},
        "formula_known_answer_tests": formula, "withheld_tests": withheld, "error_tests": errors,
        "contract_assertions": [
            {"assertion_id": "HRE_UNIT_OWN_VALUE", "expected": "each unit DR is a fraction of that unit's own value; building level = sum of DR x share"},
            {"assertion_id": "HRE_T4_UNIT_NEEDS_ACKNOWLEDGEMENT", "expected": "walls and windows answer only with the screening acknowledgement"},
            {"assertion_id": "HRE_NO_DOLLARS", "expected": "no scenario dollars until the consumer qualifies the building value"},
            {"assertion_id": "HRE_WITHHELD_NOT_ZERO", "expected": "exposed equipment and contents are withheld, never zero"},
        ],
    }


def workbook(art: dict, kats: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for line in [
        [f"{CELL_ID} model v0.1 / docs r1 — noncanonical proposal (review workbook)"],
        ["Generated by scripts/reference_helpers/build_hail_real_estate_v0_1_package.py; do not edit by hand."],
        ["Damage ratio = fraction of the failure unit's own replacement value; axis = Hazard MESH75 hail diameter, mm."],
        ["Building level = unit DR x unit share of structure replacement cost new (see failure_unit_shares)."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True)

    def sheet(title, header, data):
        s = wb.create_sheet(title)
        s.append(header)
        for c in s[1]:
            c.font = Font(bold=True)
        for row in data:
            s.append(row)

    recs = [r for p in art["pathways"] for r in p["curve_records"]]
    sheet("curve_records", ["curve_id", "failure_unit_id", "reference_building_id", "mesh_diameter_mm", "damage_ratio"],
          [[r["curve_id"], r["failure_unit_id"], r["selector_match"]["reference_building_id"], x, y]
           for r in recs for x, y in r["parameters"]["points"]])
    sheet("failure_unit_shares", ["failure_unit_id", "reference_building_id", "reference_profile", "share_of_structure_rcn"],
          [[u["id"], b, BUILDINGS[b]["profile"], s] for u in art["failure_units"] if "share_of_structure_rcn" in u
           for b, s in u["share_of_structure_rcn"].items()])
    for name, fname in [("component_fragility", f"COMPONENT_FRAGILITY_TABLE_{TAG}.csv"),
                        ("sources", f"SOURCE_REGISTER_{TAG}.csv"), ("claims", f"CLAIM_PARAMETER_REGISTER_{TAG}.csv"),
                        ("parameter_tiers", f"PARAMETER_TIER_TABLE_{TAG}.csv"), ("value_crosswalk", f"VALUE_CROSSWALK_{TAG}.csv"),
                        ("sensitivity", f"SENSITIVITY_{TAG}.csv"), ("nri_check", f"NRI_CHECK_{TAG}.csv")]:
        data = rows(fname)
        sheet(name, list(data[0]), [list(r.values()) for r in data])
    sheet("known_answer_tests", ["test_id", "kind", "expected"],
          [[t["test_id"], k, json.dumps(t.get("expected", t.get("expected_error_code")))]
           for k in ("formula_known_answer_tests", "withheld_tests", "error_tests") for t in kats[k]])
    for s in wb.worksheets:
        s.column_dimensions["A"].width = 42
    wb.properties.creator = f"{CELL_ID} builder"
    fixed = datetime.datetime(2026, 10, 18)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.save(path)


def main() -> None:
    pts = ordinates()
    art = artifact(pts)
    kats = known_answer_tests(pts)
    out = ROOT / CELL
    (out / f"{TAG}__curve_artifact.json").write_text(json.dumps(art, indent=2) + "\n")
    (out / f"{TAG}__capability.json").write_text(json.dumps(art["capability_declaration"], indent=2) + "\n")
    (out / f"known_answer_tests_{TAG}.json").write_text(json.dumps(kats, indent=2) + "\n")
    workbook(art, kats, out / f"damage_curve_records_{TAG}.xlsx")
    n = sum(len(kats[k]) for k in ("formula_known_answer_tests", "withheld_tests", "error_tests"))
    nrec = sum(len(p["curve_records"]) for p in art["pathways"])
    print(f"wrote artifact ({nrec} curve records), capability, {n} known-answer tests, workbook")


if __name__ == "__main__":
    main()
