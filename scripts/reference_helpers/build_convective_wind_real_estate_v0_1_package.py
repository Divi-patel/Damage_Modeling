"""Deterministic builder for the noncanonical convective wind × real estate model-v0.1 / docs-r1 package.

Two pathways (straight_line_convective, tornado_direct_hit) × the three Hazard reference profiles r1, from the EF scale
degree-of-damage tables priced through the FEMA Hazus damage states (decision log D-CWRE-01 to D-CWRE-16).

Reads the governed CSVs in docs/cells/convective_wind_real_estate/proposed/ (the ordinates come from
derive_convective_wind_real_estate_ef_curves.py) and writes:
    convective_wind_real_estate__model_v0_1__docs_r1__curve_artifact.json   (bundle v3, embedded capability)
    convective_wind_real_estate__model_v0_1__docs_r1__capability.json       (identical standalone capability)
    known_answer_tests_convective_wind_real_estate__model_v0_1__docs_r1.json
    damage_curve_records_convective_wind_real_estate__model_v0_1__docs_r1.xlsx

Known-answer expectations are computed here from the ordinates CSV with a separate, minimal interpolation, not by
calling the evaluator. Requires openpyxl for the workbook.

Usage (from the repository root):
    python scripts/reference_helpers/build_convective_wind_real_estate_v0_1_package.py
"""
from __future__ import annotations

import csv
import datetime
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CELL_ID = "convective_wind_real_estate"
CELL = f"docs/cells/{CELL_ID}/proposed"
TAG = f"{CELL_ID}__model_v0_1__docs_r1"
SL, TOR = "straight_line_convective", "tornado_direct_hit"
UNIT = "CWRE_BUILDING_STRUCTURE"
CONTENTS = "CWRE_BUILDING_CONTENTS"
AXES = {SL: "convective_gust_3s_10m_mph", TOR: "tornado_gust_3s_10m_mph"}
LO, HI = 0.0, 255.0
ONSET = 96.0
EXPECTED_RECORDS = 6
BUILDINGS = {
    "RE_RES_SFD_1S_SLAB": {"profile": "re_house_1s_wood_slab", "di": "FR12", "di_number": 2,
                           "flags": ["EF_FR12_ONE_AND_TWO_FAMILY_RESIDENCE", "HOUSE_FOUNDATION_SHIFT_COST_JUDGMENT"]},
    "RE_IND_WAREHOUSE": {"profile": "re_warehouse_1s_steel_slab", "di": "MBS", "di_number": 21,
                         "flags": ["EF_MBS_METAL_BUILDING_SYSTEM", "MBS_INDICATOR_INCLUDES_LARGE_OVERHEAD_DOORS"]},
    "RE_COM_OFFICE_3S_SLAB": {"profile": "re_office_3s_slab", "di": "LRB", "di_number": 17,
                              "flags": ["EF_LRB_LOW_RISE_1_TO_4_STOREYS"]},
}
PATHWAY_FLAGS = {
    SL: ["STRAIGHT_LINE_T4_TRANSFER_OF_TORNADO_EF_EVIDENCE", "NO_DAMAGE_BELOW_96_MPH_ONSET",
         "ONSET_FIXED_NOT_LOCATION_SPECIFIC"],
    TOR: ["TORNADO_EF_DEGREE_OF_DAMAGE_EVIDENCE", "TORNADO_SPECIFIC_EFFECTS_INSIDE_EF_OBSERVATIONS",
          "EF_MIDPOINT_GUST_ACCEPTED_AS_T4_BRIDGE"],
}
ALWAYS = ["NONCANONICAL_PROPOSAL", "SCREENING_EF_EXPERT_ELICITED_DEGREES_OF_DAMAGE", "WHOLE_BUILDING_SOURCE_ATOM",
          "COST_RATIOS_ASSIGNED_THROUGH_HAZUS_DAMAGE_STATES", "EF_GUST_AT_POINT_OF_DAMAGE_NOT_STATED_EXPOSURE_C",
          "REFERENCE_BUILDING_ASSUMPTIONS_NOT_OBSERVED", "NOT_FIELD_OR_CLAIMS_CALIBRATED",
          "PLATFORM_CODES_PROPOSED_NOT_ISSUED"]


def rows(name: str) -> list[dict[str, str]]:
    with open(ROOT / CELL / name, newline="") as fh:
        return list(csv.DictReader(fh))


def ordinates() -> dict[tuple[str, str], list[list[float]]]:
    pts: dict[tuple[str, str], list[list[float]]] = {}
    for r in rows(f"CURVE_ORDINATES_{TAG}.csv"):
        pts.setdefault((r["reference_building_id"], r["pathway_id"]), []).append(
            [float(r["gust_3s_10m_mph"]), float(r["damage_ratio"])])
    return {k: sorted(v) for k, v in pts.items()}


def curve_ids() -> dict[tuple[str, str], str]:
    return {(r["reference_building_id"], r["pathway_id"]): r["curve_id"] for r in rows(f"CURVE_ORDINATES_{TAG}.csv")}


def interp(points: list[list[float]], x: float) -> float:
    i = max(j for j, p in enumerate(points) if p[0] <= x)
    if points[i][0] == x:
        return points[i][1]
    (x0, y0), (x1, y1) = points[i], points[i + 1]
    return y0 + (y1 - y0) * (x - x0) / (x1 - x0)


def curve_records(pts, pathway: str) -> list[dict]:
    ids = curve_ids()
    out = []
    for building, spec in BUILDINGS.items():
        p = pts[(building, pathway)]
        onset = (f"; straight-line onset {ONSET:g} mph: DR_SL = max(0, (DR - DR({ONSET:g})) / (1 - DR({ONSET:g})))"
                 if pathway == SL else "; no onset")
        out.append({
            "curve_id": ids[(building, pathway)], "pathway_id": pathway, "failure_unit_id": UNIT,
            "curve_form": "piecewise_linear", "x_axis": AXES[pathway],
            "y_axis": "failure_unit_damage_ratio", "parameters": {"points": p},
            "valid_range": [p[0][0], p[-1][0]],
            "interpolation_policy": "linear_between_source_knots",
            "extrapolation_policy": "knots every 1 mph from 0 to 255 mph (tabulated from the EF-based expression" + onset
                                    + "); above 255 mph withholds; negative gust rejected; no clamp",
            "selector_match": {"reference_building_id": building},
            "source_parameter_refs": [
                f"CWRE-S001#Appendix_A:{spec['di']}",
                f"CWRE-S002#efscale/{spec['di_number']}.html",
                "CWRE-S003#Section_8.1.4.3",
                f"CWRE-S003#{'Table_5-44' if spec['di'] == 'FR12' else 'Table_5-74' if spec['di'] == 'MBS' else 'Table_5-78'}"],
            "metadata_flags": spec["flags"] + PATHWAY_FLAGS[pathway],
        })
    return out


def capability() -> dict:
    def pathway_cap(pathway: str) -> dict:
        return {
            "pathway_id": pathway,
            "failure_unit_scalar_dr": "conditional",
            "scenario_loss_given_value_basis": "withheld",
            "curve_intrinsic_spread": "not_carried",
            "populated_emit_modes": ["scalar_mean"],
            "conditions": [
                "reference_building_id must be one of the three Hazard reference profiles r1 buildings",
                f"{AXES[pathway]} is the 3-second gust at 10 m in open terrain, mph; no terrain or height adjustment in Damage",
                "records are valid from 0 to 255 mph; above 255 withholds; a negative gust is rejected",
                ("straight-line damage is zero at or below 96 mph (fixed onset, not location-specific)" if pathway == SL
                 else "an EF rating is not an input; Hazard's EF band midpoint gust is accepted as a damage-inferred gust"),
                ("straight-line records are a T4 transfer of tornado-derived EF evidence" if pathway == SL
                 else "tornado-specific effects (pressure drop, rotation, debris) are inside the EF observations; no separate add-on"),
                "scenario dollars stay withheld until the consumer qualifies the building value",
                "owner acceptance of the scope and selections is required before any canonical use",
            ],
            "limitation_flags": ALWAYS + PATHWAY_FLAGS[pathway] + ["VALUE_COVERAGE_NOT_ITEMIZED_BY_SOURCE",
                                                                   "MULTI_PERIL_SAME_BUILDING_COMBINATION_CONSUMER_OWNED"],
            "withheld_failure_units": [{"failure_unit_id": CONTENTS, "reason_codes": [
                "NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT", "CONTENTS_DEFERRED_SEPARATE_FAILURE_UNIT"]}],
        }

    return {
        "schema_version": "capability_declaration.v3", "cell_id": CELL_ID,
        "canonical_runtime_artifact": False,
        "pathway_capabilities": [pathway_cap(SL), pathway_cap(TOR)],
        "consumer_annual_metrics": {
            "computation_owner": "downstream_consumer",
            "status_before_promotion": "withheld_noncanonical_proposal",
            "status_after_promotion": "withheld",
            "prerequisites": ["owner acceptance of scope and curve selections",
                              "consumer-qualified replacement cost new for each reference building",
                              "Hazard straight-line frequency and tornado hit geometry reviewed (NRI check)",
                              "exact model/docs/schema/SHA pin"],
            "limitation_flags": ["ANNUAL_AND_TAIL_METRICS_CONSUMER_OWNED", "VALUE_COVERAGE_NOT_ITEMIZED_BY_SOURCE",
                                 "NRI_CHECK_STRAIGHT_LINE_ABOVE_TORNADO_BELOW"],
        },
        "cap_binding": {
            "policy": "consumer_enforced_fail_closed", "enforcement_owner": "downstream_consumer",
            "checks_required": ["exact model/docs/schema/SHA pin",
                                "record routing by pathway_id and reference_building_id",
                                "each pathway reads only its own gust field",
                                "a tornado-spawned or hurricane gust is routed to one pathway only (no double count)",
                                "all perils on one building in one event together at most 100% of structure value"],
            "action_if_fail": "withhold the affected DR or downstream metric; do not substitute another record or clamp",
        },
        "promotion_gate": {"status": "blocked", "required_before_canonical_use": [
            "owner accepts D-CWRE-01 to D-CWRE-16",
            "consumer value qualification of each reference budget (shared with flood_real_estate)",
            "reviewer accepts the T4 straight-line transfer and the cost assignment",
            "independent review of the package"]},
    }


def axis(pathway: str) -> dict:
    common = {"unit": "mph", "valid_range": [LO, HI], "averaging_period_s": 3, "reference_height_m": 10,
              "terrain_basis": "open terrain (ASCE Exposure C) as delivered by Hazard; EF speeds are point-of-damage estimates",
              "permitted_proxy_fields": [],
              "extrapolation_policy": "gust < 0 rejects; 0-255 mph straight lines between 1-mph knots; > 255 withholds"}
    if pathway == SL:
        return {"id": "CONVECTIVE_GUST_3S_10M_OPEN_TERRAIN_MPH", "preferred_input_field": AXES[SL], **common,
                "bridge": "none in Damage; Hazard delivers the ASCE 7-22 3-s gust at 10 m, Exposure C",
                "category_guardrail": "Event labels (severe-wind report, derecho) are not damage inputs."}
    return {"id": "TORNADO_GUST_3S_10M_MPH", "preferred_input_field": AXES[TOR], **common,
            "bridge": "none in Damage; Hazard's EF band midpoint gust is accepted as a damage-inferred T4 bridge",
            "category_guardrail": "An EF rating is context only and is never evaluated as a speed."}


def pathway_block(pts, pathway: str) -> dict:
    if pathway == SL:
        scope = {"included": ["thunderstorm straight-line wind on the building: downburst, gust front, derecho outflow"],
                 "excluded": ["tornado (tornado_direct_hit)", "tropical-cyclone wind (tropical_cyclone_wind_real_estate)",
                              "hail (future hail × real estate cell)", "lightning, rainfall flooding",
                              "contents, loss of use, business interruption, land, site works"],
                 "deferred_pathways": [],
                 "compound_event_rule": "Route each gust to one pathway; hail and wind on one building in one storm together at most 100% of structure value (consumer)."}
    else:
        scope = {"included": ["tornado direct hit on the building after Hazard resolves the hit: wind loading with the "
                              "pressure, rotation and debris effects present in the surveyed EF damage"],
                 "excluded": ["tornado frequency and path probability (Hazard)", "straight-line wind (straight_line_convective)",
                              "tropical-cyclone wind", "hail", "contents, loss of use, business interruption"],
                 "deferred_pathways": [],
                 "compound_event_rule": "A tornado spawned by a tropical cyclone is priced here or in tropical_cyclone_wind_real_estate, never both (consumer partition)."}
    return {
        "pathway_id": pathway,
        "hazard_scope": scope,
        "hazard_axis": axis(pathway),
        "reference_profile_binding": {b: s["profile"] for b, s in BUILDINGS.items()},
        "record_routing_fields": ["reference_building_id"],
        "selector_logic": [{"field": "reference_building_id", "required": True, "default": None,
                            "allowed": list(BUILDINGS),
                            "effect": "Fixed identity of the reference building; picks the EF damage indicator."}],
        "conditioner_logic": [],
        "ef_damage_indicators": {b: {"di": s["di"], "di_number": s["di_number"]} for b, s in BUILDINGS.items()},
        "exposure_contract": {
            "required_subject_grain": "one reference building",
            "value_rule": "Scenario dollars need structure replacement cost new for the same building; the consumer qualifies the reference budget.",
            "terrain_rule": "No terrain class: EF speeds have none. The open-terrain gust is used as delivered.",
        },
        "failure_unit_coverage": [
            {"failure_unit_id": UNIT, "status": "conditional_numeric_screening_curve", "reason_codes": []},
            {"failure_unit_id": CONTENTS, "status": "withheld_not_zero",
             "reason_codes": ["NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT", "CONTENTS_DEFERRED_SEPARATE_FAILURE_UNIT"]},
        ],
        "curve_records": curve_records(pts, pathway),
    }


def artifact(pts) -> dict:
    return {
        "schema_version": "damage_curve_record_bundle.v3", "schema_status": "proposed_draft",
        "cell_id": CELL_ID, "damage_code_id": "CONVECTIVE_WIND_REAL_ESTATE_EF_SCREENING_V0_1",
        "semantic_damage_model_version": "model v0.1", "documentation_revision": "docs r1",
        "lifecycle_state": "scaffold", "promotion_status": "proposed", "review_status": "not_yet_reviewed",
        "model_grade": "screening_expert_elicited_damage_degrees_priced_by_damage_state",
        "package_release": "unreleased", "package_baseline": "library v2.5", "package_inclusion_status": "not_included",
        "canonical_runtime_artifact": False,
        "source_dossier": f"{CELL}/{CELL_ID}_curve_derivation_dossier__model_v0_1__docs_r1.md",
        "source_workbook": f"{CELL}/damage_curve_records_{TAG}.xlsx",
        "known_answer_tests": f"{CELL}/known_answer_tests_{TAG}.json",
        "source_register": f"{CELL}/SOURCE_REGISTER_{TAG}.csv",
        "claim_parameter_register": f"{CELL}/CLAIM_PARAMETER_REGISTER_{TAG}.csv",
        "value_crosswalk": f"{CELL}/VALUE_CROSSWALK_{TAG}.csv",
        "failure_units": [
            {"id": UNIT, "subsystem": "BUILDING", "component": "BUILDING_STRUCTURE_ASSEMBLY",
             "treatment": "primary_nonzero", "y_axis": "failure_unit_damage_ratio",
             "denominator": "structure replacement cost new of the same reference building (no land, contents, site works or depreciation)",
             "platform_code_status": "proposed_awaiting_platform_codes",
             "source_atom_note": "Whole-building EF damage indicator priced through whole-building Hazus damage states; not an aggregate of component records."},
            {"id": CONTENTS, "subsystem": "BUILDING", "component": "CONTENTS", "treatment": "withheld",
             "y_axis": "failure_unit_damage_ratio", "denominator": "contents value (deferred)",
             "withheld_reason_codes": ["NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT", "CONTENTS_DEFERRED_SEPARATE_FAILURE_UNIT"],
             "platform_code_status": "proposed_awaiting_platform_codes"},
        ],
        "pathways": [pathway_block(pts, SL), pathway_block(pts, TOR)],
        "value_linkage": {
            "curve_denominator": "structure replacement cost new of the reference building",
            "scenario_loss_formula": "loss = DR * structure_replacement_cost_new_usd",
            "scenario_loss_status": "withheld_consumer_value_qualification",
            "rejected_denominators": ["depreciated value", "market value", "assessed value"],
            "implicit_default_profile": None,
        },
        "parameter_tier_table": rows(f"PARAMETER_TIER_TABLE_{TAG}.csv"),
        "evaluation_contract": {
            "supported_pathway_ids": [SL, TOR], "supported_failure_unit_id": UNIT,
            "required_selector_fields": ["reference_building_id"],
            "required_conditioner_fields": [],
            "accepted_axis_payloads": [{"pathway_id": p, "mode": "gust_mph", "required_fields": [AXES[p]]} for p in (SL, TOR)],
            "metadata_flags_always": ALWAYS,
            "out_of_range_behavior": "above 255 mph: withhold; negative: error",
            "reference_evaluator": "scripts/reference_helpers/convective_wind_real_estate_curve_eval.py",
        },
        "emit_contract": {
            "schema_version": "damage_emit.v2",
            "populated_emit_modes_for_this_cell": ["scalar_mean"],
            "supported_output": "conditional scalar DR for CWRE_BUILDING_STRUCTURE, per pathway",
            "scenario_loss": "withheld until the consumer qualifies the building value",
            "annual_and_tail_metrics": "withheld; consumer-owned",
            "withheld_units_emit": "explicit null DR plus reason codes; never zero",
        },
        "capability_declaration": capability(),
        "derivation_rationale": {
            "source_spine": "EF scale degrees of damage (Texas Tech WSEC 2006 Rev 2; NOAA SPC) for FR12, MBS and LRB, priced through the FEMA Hazus damage states (Hazus 7.0 Hurricane TM §8.1.4.3 and Tables 5-44, 5-74, 5-78).",
            "adopted_form": "expected cost over ordered degrees of damage, each lognormal in gust (median = expected speed, beta from the bounds read as a 90% range); tabulated every 1 mph; straight-line onset 96 mph",
            "why_not_hurricane_curves": "Separate evidence per hazard (owner choice, D-CWRE-03); hurricane-curve reuse is the documented alternative.",
            "parameter_tier_mix": "Speeds T3 (expert elicitation); cost assignment, spread reading, onset and the straight-line transfer T4; range T2.",
            "nri_check": "Straight-line yearly loss 13-35x NRI with the onset (370-500x without); tornado 0.11-0.26x NRI (hazard-side geometry).",
        },
        "legacy_comparison": {"prior_cell_state": "none (new cell)",
                              "changes_from_prior": [],
                              "legacy_arrays_used": "none",
                              "canonical_or_consumer_change": False},
    }


def known_answer_tests(pts) -> dict:
    ids = curve_ids()

    def req(building, pathway, gust, **extra):
        return {"pathway_id": pathway, "failure_unit_id": UNIT, "reference_building_id": building, AXES[pathway]: gust, **extra}

    formula, withheld, errors = [], [], []
    for building in BUILDINGS:
        for pathway in (SL, TOR):
            p, cid = pts[(building, pathway)], ids[(building, pathway)]
            for x in (0, 60, 80, 96, 100, 110, 120, 130, 140, 150, 160, 180, 200, 220, 255):
                formula.append({"test_id": f"{cid}_KNOT_{x}MPH", "input": req(building, pathway, float(x)),
                                "expected": {"status": "conditional", "curve_id": cid, "failure_unit_damage_ratio": interp(p, float(x))}})
            for x in (75.0, 98.0, 123.0, 150.5, 183.0, 227.0, 96.5, 252.8):  # EF midpoints; just above onset; Hazard bound
                formula.append({"test_id": f"{cid}_AT_{x:g}MPH", "input": req(building, pathway, x),
                                "expected": {"status": "conditional", "curve_id": cid, "failure_unit_damage_ratio": interp(p, x)}})
            withheld.append({"test_id": f"{cid}_ABOVE_RANGE", "input": req(building, pathway, HI + 0.0001),
                             "expected": {"status": "withheld", "reason_code": "ABOVE_SOURCE_VALID_RANGE"}})
            withheld.append({"test_id": f"{cid}_CONTENTS_WITHHELD", "input": {**req(building, pathway, 130.0), "failure_unit_id": CONTENTS},
                             "expected": {"status": "withheld", "reason_code": "NO_RUNTIME_CURVE_FOR_PATHWAY_UNIT"}})
        formula.append({"test_id": f"{building}_STRAIGHT_LINE_ZERO_AT_ONSET", "input": req(building, SL, ONSET),
                        "expected": {"status": "conditional", "curve_id": ids[(building, SL)], "failure_unit_damage_ratio": 0.0,
                                     "metadata_flags_include": ["NO_DAMAGE_BELOW_96_MPH_ONSET"]}})
    base = req("RE_RES_SFD_1S_SLAB", SL, 130.0)
    drop = lambda k: {key: v for key, v in base.items() if key != k}  # noqa: E731
    for tid, inp, code in [
        ("PATHWAY_REQUIRED", drop("pathway_id"), "PATHWAY_ID_REQUIRED"),
        ("PATHWAY_UNKNOWN", {**base, "pathway_id": "tropical_cyclone_wind"}, "PATHWAY_ID_UNKNOWN"),
        ("BUILDING_REQUIRED", drop("reference_building_id"), "REFERENCE_BUILDING_ID_REQUIRED"),
        ("BUILDING_UNKNOWN", {**base, "reference_building_id": "RE_COM_LOWRISE"}, "REFERENCE_BUILDING_ID_UNKNOWN"),
        ("BUILDING_PROFILE_ID_NOT_A_SELECTOR", {**base, "reference_building_id": "re_office_3s_slab"}, "REFERENCE_BUILDING_ID_UNKNOWN"),
        ("AXIS_REQUIRED", drop(AXES[SL]), "AXIS_PAYLOAD_REQUIRED"),
        ("AXIS_OF_OTHER_PATHWAY", {**drop(AXES[SL]), AXES[TOR]: 130.0}, "AXIS_PAYLOAD_REQUIRED"),
        ("TORNADO_EF_RATING_NOT_A_GUST", {"pathway_id": TOR, "failure_unit_id": UNIT,
                                          "reference_building_id": "RE_RES_SFD_1S_SLAB", "ef_rating": "EF2"}, "AXIS_PAYLOAD_REQUIRED"),
        ("AXIS_NOT_NUMBER", {**base, AXES[SL]: "130"}, "AXIS_VALUE_INVALID"),
        ("AXIS_BOOLEAN", {**base, AXES[SL]: True}, "AXIS_VALUE_INVALID"),
        ("AXIS_NEGATIVE", {**base, AXES[SL]: -5.0}, "AXIS_OUTSIDE_VALID_RANGE"),
        ("UNIT_REQUIRED", drop("failure_unit_id"), "FAILURE_UNIT_ID_REQUIRED"),
        ("UNIT_UNKNOWN", {**base, "failure_unit_id": "CWRE_ROOF"}, "FAILURE_UNIT_ID_UNKNOWN"),
    ]:
        errors.append({"test_id": tid, "input": inp, "expected_error_code": code})
    return {
        "schema_version": "known_answer_tests.v3", "cell_id": CELL_ID,
        "semantic_damage_model_version": "model v0.1", "documentation_revision": "docs r1",
        "artifact_schema_version": "damage_curve_record_bundle.v3", "artifact_status": "noncanonical_proposal",
        "absolute_tolerance": 1e-12,
        "equation": {"runtime": "straight lines between the 1-mph knots of the selected record",
                     "axis": "3-second gust at 10 m, open terrain, mph; each pathway reads its own field",
                     "ranges": "0 to 255 mph evaluates; above 255 mph withholds; negative is an error"},
        "formula_known_answer_tests": formula, "withheld_tests": withheld, "error_tests": errors,
        "contract_assertions": [
            {"assertion_id": "CWRE_TWO_GOVERNED_PATHWAYS", "expected": "tornado and straight-line wind are separate pathways with separate records, not a shifted copy"},
            {"assertion_id": "CWRE_EF_RATING_NEVER_EVALUATED", "expected": "an EF rating without a gust is an error"},
            {"assertion_id": "CWRE_NO_DOLLARS", "expected": "no scenario dollars until the consumer qualifies the building value"},
            {"assertion_id": "CWRE_CONTENTS_WITHHELD_NOT_ZERO", "expected": "contents are withheld, never reported as zero"},
        ],
    }


def workbook(art: dict, kats: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for line in [
        [f"{CELL_ID} model v0.1 / docs r1 — noncanonical proposal (review workbook)"],
        ["Generated by scripts/reference_helpers/build_convective_wind_real_estate_v0_1_package.py; do not edit by hand."],
        ["Damage ratio = fraction of structure replacement cost new; axis = 3-s gust at 10 m, open terrain, mph."],
        ["Straight lines between 1-mph knots; 0-255 mph; straight-line onset 96 mph; tornado no onset."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True)

    def sheet(title, header, data):
        s = wb.create_sheet(title)
        s.append(header)
        for c in s[1]:
            c.font = Font(bold=True)
        for row in data:
            s.append(row)

    recs = [r for p in art["pathways"] for r in p["curve_records"]]
    sheet("curve_records", ["curve_id", "pathway_id", "reference_building_id", "gust_mph", "damage_ratio"],
          [[r["curve_id"], r["pathway_id"], r["selector_match"]["reference_building_id"], x, y]
           for r in recs for x, y in r["parameters"]["points"]])
    sheet("reference_buildings", ["reference_building_id", "reference_profile", "ef_damage_indicator", "ef_di_number"],
          [[b, s["profile"], s["di"], s["di_number"]] for b, s in BUILDINGS.items()])
    for name, fname in [("ef_degrees_of_damage", f"EF_DEGREE_OF_DAMAGE_TABLE_{TAG}.csv"),
                        ("sources", f"SOURCE_REGISTER_{TAG}.csv"), ("claims", f"CLAIM_PARAMETER_REGISTER_{TAG}.csv"),
                        ("parameter_tiers", f"PARAMETER_TIER_TABLE_{TAG}.csv"), ("value_crosswalk", f"VALUE_CROSSWALK_{TAG}.csv"),
                        ("sensitivity", f"SENSITIVITY_{TAG}.csv"), ("nri_check", f"NRI_CHECK_{TAG}.csv")]:
        data = rows(fname)
        sheet(name, list(data[0]), [list(r.values()) for r in data])
    sheet("known_answer_tests", ["test_id", "kind", "expected"],
          [[t["test_id"], k, json.dumps(t.get("expected", t.get("expected_error_code")))]
           for k in ("formula_known_answer_tests", "withheld_tests", "error_tests") for t in kats[k]])
    for s in wb.worksheets:
        s.column_dimensions["A"].width = 42
    wb.properties.creator = f"{CELL_ID} builder"
    fixed = datetime.datetime(2026, 10, 18)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.save(path)


def main() -> None:
    pts = ordinates()
    art = artifact(pts)
    kats = known_answer_tests(pts)
    out = ROOT / CELL
    (out / f"{TAG}__curve_artifact.json").write_text(json.dumps(art, indent=2) + "\n")
    (out / f"{TAG}__capability.json").write_text(json.dumps(art["capability_declaration"], indent=2) + "\n")
    (out / f"known_answer_tests_{TAG}.json").write_text(json.dumps(kats, indent=2) + "\n")
    workbook(art, kats, out / f"damage_curve_records_{TAG}.xlsx")
    n = sum(len(kats[k]) for k in ("formula_known_answer_tests", "withheld_tests", "error_tests"))
    nrec = sum(len(p["curve_records"]) for p in art["pathways"])
    print(f"wrote artifact ({nrec} curve records), capability, {n} known-answer tests, workbook")


if __name__ == "__main__":
    main()
