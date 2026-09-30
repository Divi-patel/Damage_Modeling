#!/usr/bin/env python3
"""Verify DD-G31 value binding against preserved responses and source bytes."""
from copy import deepcopy
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from build_tropical_cyclone_wind_wind_v1_3_value_binding import (
    BASIS, CELL, DEST, PARENT, POLICY, PROFILE, ROOT, SHARE, STEM, build,
)
from tropical_cyclone_wind_wind_curve_eval import (
    SUPPORTED_FAILURE_UNIT, TropicalCycloneWindEvaluationError, evaluate_damage_call,
)
from validate_tropical_cyclone_wind_wind_v1_2_tower_proxy import (
    result, validate_reproduction, validate_schemas,
)


def validate() -> dict:
    path = DEST / f"{STEM}__curve_artifact.json"
    artifact = json.loads(path.read_text())
    parent = json.loads(PARENT.read_text())
    assert artifact == build(), "generated artifact differs from source builder"
    assert artifact["pathways"][0]["curve_records"] == parent["pathways"][0]["curve_records"]
    capability = json.loads((DEST / f"{STEM}__capability.json").read_text())
    assert capability == artifact["capability_declaration"]
    kats = json.loads((CELL / "current/known_answer_tests_tropical_cyclone_wind_wind__model_v1_2__docs_r2.json").read_text())
    reproduction = validate_reproduction(parent, artifact, kats)
    migrated = deepcopy(kats)
    replacements = {
        "TCWW_OWNER_APPROVED_3P3MW_FOR_CANONICAL_5MW_TOWER_ONLY_V1": POLICY,
        "CONUS_WIND_FARM_REFERENCE_V1": PROFILE,
        "CONUS_WIND_FARM_TOWER_16PCT_V1": BASIS,
    }
    for group in ("proxy_known_answer_tests", "negative_contract_tests"):
        for case in migrated[group]:
            case["input"] = {k: replacements.get(v, v) if isinstance(v, str) else v
                             for k, v in case["input"].items()}
    for case in migrated["proxy_known_answer_tests"]:
        emit = evaluate_damage_call(artifact, case["input"])
        answer = result(emit, SUPPORTED_FAILURE_UNIT)
        assert math.isclose(answer["scalar_central_dr"], case["expected"]["failure_unit_damage_ratio"], abs_tol=1e-12)
        assert emit["exposure_used"]["covered_value_share_of_project_tiv"] == SHARE
        assert emit["exposure_used"]["covered_subsystems"] == ["tower"]
        assert "NON_TOWER_PHYSICAL_VALUE_WITHHELD_NOT_ZERO" in answer["metadata_flags"]
    for case in migrated["negative_contract_tests"]:
        try:
            evaluate_damage_call(artifact, case["input"])
        except TropicalCycloneWindEvaluationError as exc:
            assert exc.code == case["expected_error_code"]
        else:
            raise AssertionError(case["test_id"])
    old_request = kats["proxy_known_answer_tests"][0]["input"]
    try:
        evaluate_damage_call(artifact, old_request)
    except TropicalCycloneWindEvaluationError as exc:
        assert exc.code == "OWNER_APPROVED_PROXY_IDENTITY_MISMATCH"
    else:
        raise AssertionError("old value identity accepted by new model")
    # Fine sweep includes both open transition boundaries and all source ranges.
    old_request = dict(old_request)
    new_request = dict(migrated["proxy_known_answer_tests"][0]["input"])
    for i in range(4001):
        old_request["tc_peak_gust_3s_10m_kmh"] = i / 10
        new_request["tc_peak_gust_3s_10m_kmh"] = i / 10
        old = result(evaluate_damage_call(parent, old_request), SUPPORTED_FAILURE_UNIT)
        new = result(evaluate_damage_call(artifact, new_request), SUPPORTED_FAILURE_UNIT)
        assert old["scalar_central_dr"] == new["scalar_central_dr"]
        assert 0 <= new["scalar_central_dr"] * 27_600_000 <= 27_600_000
    binding = artifact["value_linkage"]["owner_approved_proxy_value_basis"]
    workbook = ROOT / "docs/method/value_basis/solar_wind_value_breakdown.xlsx"
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == binding["source_workbook_sha256"]
    book = load_workbook(workbook, read_only=True, data_only=True)
    values = list(book["Wind_Map"].iter_rows(min_row=9, max_row=9, values_only=True))[0]
    assert "TOWER" in values and 276 in values
    book.close()
    assert math.isclose(162_300_000 * SHARE, 27_600_000, rel_tol=0, abs_tol=1e-8)
    assert 162_300_000 - 27_600_000 == binding["uncovered_value_usd_at_activation"]
    schemas = validate_schemas(artifact, capability, migrated)
    return {
        "status": "passed_proposal_contract_checks",
        "artifact_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "source_native_known_answers": reproduction,
        "proxy_known_answers": len(migrated["proxy_known_answer_tests"]),
        "negative_contracts": len(migrated["negative_contract_tests"]) + 1,
        "exact_response_sweep": 4001, "source_workbook_sha256": binding["source_workbook_sha256"],
        "tower_usd": 27_600_000, "full_tiv_usd": 162_300_000,
        "schemas": schemas, "national_consumer_gate": "separate Hazard receipt required",
        "canonical": False,
    }


if __name__ == "__main__":
    receipt = validate()
    (DEST / f"{STEM}__verification.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt, indent=2))
