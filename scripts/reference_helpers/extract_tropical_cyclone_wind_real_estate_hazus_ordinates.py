"""Extract the proposed tropical-cyclone wind × real estate curve ordinates from the FEMA Hazus wind extract.

Reads the FEMA-provided Hazus hurricane damage/loss-function extract (``HazusWindDamFunctions_Hazus61.xlsx``,
public mirror of the OS-Climate physrisk project) and writes the governed ordinates CSV for the three
reference buildings × five Hazus terrain classes (15 records, 41 points each, 50–250 mph). Optionally
cross-checks every written value against NHERI SimCenter's independent Hazus 5.1 extraction
(``loss_repair.csv`` of the Damage and Loss Model Library, "Hazus v5.1 original").

The source files are not in the repository; their SHA-256 values are recorded in the source register
(TCRE-S002, TCRE-S003) and are checked here. Requires openpyxl.

Usage (from the repository root):
    python scripts/reference_helpers/extract_tropical_cyclone_wind_real_estate_hazus_ordinates.py \
        <HazusWindDamFunctions_Hazus61.xlsx> [<simcenter loss_repair.csv>]
"""
from __future__ import annotations

import csv
import hashlib
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/cells/tropical_cyclone_wind_real_estate/proposed/CURVE_ORDINATES_tropical_cyclone_wind_real_estate__model_v0_1__docs_r1.csv"
EXTRACT_SHA256 = "0fe93bbdc5d4ae3a3eace7f8147436ad63f381175688671aebd769674d1bc353"
SIMCENTER_SHA256 = "6f023f82c4ab56e57d0d20becf72e0c46a13cf4a80ebd28c3ac50f6bb83de222"
BUILDING_LOSS = 5  # huDamLossFunDescription: DamLossDescID 5 = "Building Loss"

# reference building -> (curve-id stem, Hazus SBT, Hazus charDescription, CaseID filter or None, SimCenter ID stem)
BUILDINGS = {
    "RE_RES_SFD_1S_SLAB": ("TCRE_RES_HAZUS_WSF1", "WSF1", "rsgabswrnorda6dtnailgdnodshtnorcshg", 1,
                           "W.SF.1.gab.0.6d.tnail.no.0"),
    "RE_COM_LOWRISE": ("TCRE_COM_HAZUS_SECBL", "SECBL", "rcburwalowshtnowiddBrd100", None,
                       "S.ECB.L.bur.0.B.std.low"),
    "RE_IND_WAREHOUSE": ("TCRE_IND_HAZUS_SPMBL", "SPMBL", "dqporshtnord100", None,
                         "S.PMB.L.0.por.std"),
}
# Hazus huTerrain: TERRAINID -> (class name used by this cell, z0 m, SimCenter suffix)
TERRAINS = {
    1: ("open", 0.03, "3"),
    2: ("light_suburban", 0.15, "15"),
    3: ("suburban", 0.35, "35"),
    4: ("light_trees", 0.70, "70"),
    5: ("trees", 1.00, "100"),
}
SOURCE_IDS = "TCRE-S001;TCRE-S002;TCRE-S003"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def read_extract(path: Path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    types = {}
    for row in wb["huListOfWindBldgTypes"].iter_rows(min_row=2, values_only=True):
        wb_id, sbt, desc, _n, case = row[:5]
        types[(sbt.strip(), desc.strip(), case)] = wb_id
    wanted = {}
    for building, (_stem, sbt, desc, case, _sc) in BUILDINGS.items():
        hits = [wid for (s, d, c), wid in types.items() if s == sbt and d == desc and (case is None or c == case)]
        if len(hits) != 1:
            raise SystemExit(f"{building}: expected one wbID for {sbt} {desc} case={case}, got {hits}")
        wanted[hits[0]] = building
    sheet = wb["huDamLossFun"]
    header = next(sheet.iter_rows(max_row=1, values_only=True))
    speeds = [int(h[2:]) for h in header[3:]]
    if speeds != list(range(50, 255, 5)):
        raise SystemExit(f"unexpected wind-speed columns {speeds[:3]}…{speeds[-3:]}")
    curves = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        wb_id, terrain, desc_id = row[:3]
        if wb_id in wanted and desc_id == BUILDING_LOSS:
            curves[(wanted[wb_id], terrain)] = (wb_id, [float(v) for v in row[3:]])
    if len(curves) != len(BUILDINGS) * len(TERRAINS):
        raise SystemExit(f"expected 15 building-loss curves, found {len(curves)}")
    return speeds, curves


def read_simcenter(path: Path) -> dict[str, list[float]]:
    out = {}
    with open(path, newline="") as fh:
        for r in csv.DictReader(fh):
            ys, xs = r["LossFunction-Theta_0"].split("|")
            if [float(x) for x in xs.split(",")] != [float(s) for s in range(50, 255, 5)]:
                continue
            out[r["ID"].removesuffix("-Cost")] = [float(y) for y in ys.split(",")]
    return out


def main(argv: list[str]) -> int:
    if len(argv) not in (2, 3):
        print(__doc__)
        return 2
    extract = Path(argv[1])
    if sha256(extract) != EXTRACT_SHA256:
        raise SystemExit("extract SHA-256 differs from TCRE-S002")
    speeds, curves = read_extract(extract)
    simcenter = None
    if len(argv) == 3:
        if sha256(Path(argv[2])) != SIMCENTER_SHA256:
            raise SystemExit("SimCenter file SHA-256 differs from TCRE-S003")
        simcenter = read_simcenter(Path(argv[2]))

    rows, max_diff = [], 0.0
    for building, (stem, sbt, desc, _case, sc_stem) in BUILDINGS.items():
        for tid, (tname, z0, sc_suffix) in TERRAINS.items():
            wb_id, values = curves[(building, tid)]
            if simcenter is not None:
                ref = simcenter[f"{sc_stem}.{sc_suffix}"]
                max_diff = max(max_diff, max(abs(a - b) for a, b in zip(values, ref)))
            for mph, value in zip(speeds, values):
                rows.append([f"{stem}_{tname.upper()}", building, sbt, desc, wb_id, tid, tname, f"{z0:.2f}",
                             mph, repr(value), SOURCE_IDS])
    if simcenter is not None and max_diff != 0.0:
        raise SystemExit(f"extract and SimCenter copy differ (max abs diff {max_diff})")

    with open(OUT, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["curve_id", "reference_building_id", "hazus_sbt", "hazus_wind_building_characteristics",
                    "hazus_wbid", "hazus_terrain_id", "hazus_terrain_class", "surface_roughness_z0_m",
                    "peak_gust_3s_10m_open_terrain_mph", "building_loss_ratio", "source_ids"])
        w.writerows(rows)
    note = "cross-checked against SimCenter: identical" if simcenter is not None else "not cross-checked"
    print(f"wrote {len(rows)} ordinates ({len(rows) // len(speeds)} records) to {OUT.relative_to(ROOT)}; {note}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
